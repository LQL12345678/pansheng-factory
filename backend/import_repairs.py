import sqlite3
from openpyxl import load_workbook
from datetime import datetime

# 数据库路径
DB_PATH = r'C:\Users\Admin\WorkBuddy\2026-05-10-task-1\backend\db\equipment.db'
EXCEL_PATH = r'C:\Users\Admin\Desktop\维修工单导出.xlsx'

# 连接数据库
conn = sqlite3.connect(DB_PATH)
cursor = conn.cursor()

# 1. 读取所有设备，建立设备名称到ID的映射
cursor.execute('SELECT id, name FROM devices')
devices = {row[1]: row[0] for row in cursor.fetchall()}
print(f'已加载 {len(devices)} 个设备')

# 2. 读取Excel文件
wb = load_workbook(EXCEL_PATH, data_only=True)
ws = wb.active

# 列索引（0-based）
COL_WORK_ORDER = 0
COL_DEVICE_NAME = 2
COL_LINE = 3
COL_REPORTER = 4
COL_AREA = 5
COL_FAULT_DESC = 6
COL_FAULT_CAUSE = 8
COL_SOLUTION = 9
COL_REPAIRMAN = 10
COL_REPORT_DATE = 11
COL_FINISH_TIME = 12
COL_DURATION = 13
COL_STOP_DURATION = 14
COL_STATUS = 15
COL_REMARKS = 16

# 特殊名称映射表
NAME_MAPPING = {
    'A线后侧段上层': 'A线后侧段回流层',  # Excel名称 -> 数据库名称
    'B线下机箱提升机': 'B线1F下机箱提升机',
    'A线下机箱提升机': 'A线1F下机箱提升机',
    'VIP线皮带线1': 'VIP皮带线1',
    'B线后测段回流层': 'B线后侧段回流层',
    'B线上机箱皮带线1': 'B线1F上机箱皮带线1',
    'B线滚筒平移机': 'B线包装滚筒平移机',
    'A线滚筒平移机': 'A线包装滚筒平移机',
}

def normalize_name(name):
    """标准化设备名称"""
    if not name:
        return None
    # 先检查映射表
    if name in NAME_MAPPING:
        return NAME_MAPPING[name]
    return name

def find_device_id(device_name, devices_dict):
    """查找设备ID，支持模糊匹配"""
    # 标准化名称
    normalized = normalize_name(device_name)
    if normalized and normalized in devices_dict:
        return devices_dict[normalized]

    # 尝试模糊匹配
    for dev_name, dev_id in devices_dict.items():
        if device_name and dev_name and (
            device_name in dev_name or dev_name in device_name or
            # 处理"上层/下层"差异
            device_name.replace('上层', '').replace('下层', '') in dev_name or
            dev_name.replace('上层', '').replace('下层', '') in device_name or
            # 处理"1F/3F"前缀差异
            device_name.replace('1F', '').replace('3F', '').strip() == dev_name.replace('1F', '').replace('3F', '').strip()
        ):
            return dev_id
    return None

# 3. 清空现有维修数据
cursor.execute('DELETE FROM repairs')
print('已清空现有维修数据')

# 4. 读取Excel数据并导入
success_count = 0
skip_count = 0
not_found_devices = set()

for row_idx in range(2, ws.max_row + 1):  # 从第2行开始（跳过表头）
    row = [ws.cell(row=row_idx, column=col+1).value for col in range(17)]

    work_order_no = row[COL_WORK_ORDER]
    device_name = row[COL_DEVICE_NAME]
    line = str(row[COL_LINE]) if row[COL_LINE] else None
    reporter = row[COL_REPORTER]
    area = row[COL_AREA]
    fault_desc = row[COL_FAULT_DESC]
    fault_cause = row[COL_FAULT_CAUSE]
    solution = row[COL_SOLUTION]
    repairman = row[COL_REPAIRMAN]
    report_date = row[COL_REPORT_DATE]
    finish_time = row[COL_FINISH_TIME]
    duration = row[COL_DURATION]
    stop_duration = row[COL_STOP_DURATION]
    status = row[COL_STATUS]
    remarks = row[COL_REMARKS]

    # 跳过空行
    if not work_order_no or not fault_desc:
        skip_count += 1
        continue

    # 转换日期
    if isinstance(report_date, datetime):
        report_date_str = report_date.strftime('%Y-%m-%d')
    elif report_date:
        report_date_str = str(report_date)[:10]
    else:
        report_date_str = datetime.now().strftime('%Y-%m-%d')

    if isinstance(finish_time, datetime):
        finish_time_str = finish_time.strftime('%Y-%m-%d')
    elif finish_time:
        finish_time_str = str(finish_time)[:10]
    else:
        finish_time_str = None

    # 查找设备ID
    device_id = find_device_id(device_name, devices)
    if not device_id:
        not_found_devices.add(device_name)
        print(f'  跳过: {work_order_no} - 设备未找到: {device_name}')
        skip_count += 1
        continue

    # 转换状态
    if status == '已完结':
        status = '已完成'

    # 插入数据
    cursor.execute('''
        INSERT INTO repairs (
            work_order_no, report_date, device_id, line, reporter, area,
            fault_desc, fault_cause, solution, repairman,
            finish_time, duration_minutes, stop_duration_minutes,
            status, remarks
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        work_order_no, report_date_str, device_id, line, reporter, area,
        fault_desc, fault_cause, solution, repairman,
        finish_time_str, duration or None, stop_duration or None,
        status or '已完成', remarks
    ))
    success_count += 1

conn.commit()

print(f'\n导入完成!')
print(f'  成功: {success_count} 条')
print(f'  跳过: {skip_count} 条')

if not_found_devices:
    print(f'\n未匹配的设备名称:')
    for d in not_found_devices:
        print(f'  - {d}')

# 验证
cursor.execute('SELECT COUNT(*) FROM repairs')
total = cursor.fetchone()[0]
print(f'\n数据库中维修记录总数: {total}')

conn.close()
